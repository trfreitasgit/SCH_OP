#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Gera a planilha executiva (.xlsx) das ordens de produção selecionadas.

Dois caminhos de entrada:

    python scripts/export_excel.py                  # linhas em JSON no stdin
    python scripts/export_excel.py --ids "1,2,3"    # consulta o Supabase (CLI/manual)

O servidor usa o caminho stdin: ele já tem o cliente Supabase e manda as linhas
prontas, o que evita importar o supabase-py e fazer uma chamada de rede dentro
do Python (eram ~3,6 s por exportação).

Contrato: a ULTIMA linha do stdout precisa ser o caminho do arquivo .xlsx
gerado. Tudo que for aviso/erro vai para o stderr, porque o servidor lê apenas
a última linha do stdout para achar o arquivo e depois o remove.
"""

import argparse
import json
import os
import re
import sys
import tempfile
import uuid
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from zoneinfo import ZoneInfo

TABELA = 'ordens_producao'

# Ordem das colunas na planilha: (campo no banco, rótulo exibido).
# 'turno' não existe no banco: é derivado de dia_analise (ver classificar_turno).
COLUNAS = [
    ('numero_op', 'OP'),
    ('total_pecas', 'Peças'),
    ('data_abertura', 'Data Abertura'),
    ('material', 'Material'),
    ('situacao', 'Situação'),
    ('observacao', 'Observação'),
    ('setor', 'Setor'),
    ('prioridade', 'Prioridade'),
    ('dia_analise', 'Dia Análise'),
    ('turno', 'Turno'),
]

ROTULOS = [rotulo for _, rotulo in COLUNAS]

CAMPOS_DATA = ('data_abertura', 'dia_analise')
CAMPOS_NUMERICOS = ('total_pecas',)

# Turnos do pagamento, fonte da verdade no server.js (classificarTurno):
#   T1: 05:20 - 13:39 | T2: 13:40 - 21:59 | T3: 22:00 - 05:19 (cruza a meia-noite)
FUSO_FABRICA = 'America/Sao_Paulo'
TURNO_T1_INICIO = 5 * 60 + 20   # 05:20
TURNO_T2_INICIO = 13 * 60 + 40  # 13:40
TURNO_T3_INICIO = 22 * 60       # 22:00

FORMATO_DATA_EXIBICAO = '%d/%m/%Y'

LARGURA_MAXIMA_COLUNA = 60
PADDING_LARGURA = 2

# ARGB, igual ao usado no dashboard.
COR_HEADER = 'FF1F3864'
COR_TEXTO_HEADER = 'FFFFFFFF'
COR_BORDA = 'FFB0B7C3'

RAIZ = Path(__file__).resolve().parents[1]


# ---------------------------------------------------------------------------
# Entrada
# ---------------------------------------------------------------------------

def carregar_env():
    """Carrega o .env da raiz do projeto (mesmo arquivo usado pelo server.js)."""
    from dotenv import load_dotenv

    load_dotenv(RAIZ / '.env', override=False)


def parse_ids(texto):
    """Converte "1,2,3" em [1, 2, 3] (sem repetidos, na ordem de entrada)."""
    if texto is None:
        raise ValueError('o argumento --ids é obrigatório')

    ids = []
    for parte in str(texto).split(','):
        bruto = parte.strip()
        if bruto == '':
            raise ValueError('a lista de IDs contém um valor vazio')
        if not bruto.isdigit() or int(bruto) <= 0:
            raise ValueError('ID inválido: "%s" (só inteiros positivos)' % bruto)
        numero = int(bruto)
        if numero not in ids:
            ids.append(numero)

    if not ids:
        raise ValueError('nenhum ID informado em --ids')

    return ids


def ler_entrada_stdin(texto):
    """Lê o payload JSON que o servidor manda no stdin.

    Formato: {"ids": [1, 2], "ordens": [{...}, {...}]}
    Retorna (ids, registros).
    """
    if texto is None or str(texto).strip() == '':
        raise ValueError('nenhuma ordem recebida no stdin')

    try:
        payload = json.loads(texto)
    except (TypeError, ValueError) as erro:
        raise ValueError('payload do stdin não é um JSON válido: %s' % erro) from erro

    if not isinstance(payload, dict):
        raise ValueError('o payload do stdin deve ser um objeto JSON')

    registros = payload.get('ordens')
    if not isinstance(registros, list):
        raise ValueError('o payload do stdin precisa trazer a lista "ordens"')

    ids = payload.get('ids')
    if ids is None:
        ids = [registro.get('id') for registro in registros if isinstance(registro, dict)]
    if not isinstance(ids, list):
        raise ValueError('o campo "ids" do payload precisa ser uma lista')

    ids_limpos = []
    for identificador in ids:
        try:
            ids_limpos.append(int(identificador))
        except (TypeError, ValueError):
            continue

    return ids_limpos, [registro for registro in registros if isinstance(registro, dict)]


def criar_cliente():
    """Cria o cliente Supabase a partir das variáveis de ambiente (caminho --ids)."""
    from supabase import create_client

    url = os.getenv('SUPABASE_URL_ORDENS') or os.getenv('SUPABASE_URL')
    chave = os.getenv('SUPABASE_SERVICE_ROLE_KEY_ORDENS') or os.getenv('SUPABASE_SERVICE_ROLE_KEY')

    if not url or not chave:
        raise RuntimeError(
            'SUPABASE_URL_ORDENS e SUPABASE_SERVICE_ROLE_KEY_ORDENS precisam estar definidas.'
        )

    return create_client(url, chave)


def buscar_ordens(cliente, ids):
    """Busca os registros completos dos IDs informados."""
    if not ids:
        return []

    resposta = cliente.table(TABELA).select('*').in_('id', list(ids)).execute()
    return list(resposta.data or [])


def ids_disponiveis(cliente, limite=1):
    """Alguns IDs existentes na tabela (usado pelos testes de integração)."""
    resposta = cliente.table(TABELA).select('id').order('id').limit(limite).execute()
    return [int(linha['id']) for linha in (resposta.data or [])]


# ---------------------------------------------------------------------------
# Transformação dos dados
# ---------------------------------------------------------------------------

@lru_cache(maxsize=1)
def _fuso_fabrica():
    """Fuso da fábrica (America/Sao_Paulo)."""
    try:
        return ZoneInfo(FUSO_FABRICA)
    except Exception as erro:  # noqa: BLE001 - sem base tz o script não tem como classificar turno
        raise RuntimeError(
            'Não foi possível carregar o fuso %s (instale o pacote "tzdata").' % FUSO_FABRICA
        ) from erro


def _parse_data_texto(texto):
    """Converte o texto vindo do PostgREST em datetime (ingênuo quando não houver fuso)."""
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', texto):
        # Data pura: vale como meia-noite no fuso da fábrica, sem conversão.
        return datetime.strptime(texto, '%Y-%m-%d')

    try:
        return datetime.fromisoformat(texto.replace('Z', '+00:00'))
    except ValueError:
        pass

    for formato in ('%d-%m-%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue

    raise ValueError('não é uma data válida: "%s"' % texto)


def instante_fabrica(valor):
    """Instante de `valor` no fuso da fábrica.

    Datas/horas sem fuso são tratadas como horário local da fábrica. Levanta
    ValueError quando o valor não for uma data/hora reconhecível.
    """
    if valor is None:
        raise ValueError('sem data')
    if isinstance(valor, datetime):
        instante = valor
    elif isinstance(valor, date):
        instante = datetime(valor.year, valor.month, valor.day)
    else:
        texto = str(valor).strip()
        if texto == '':
            raise ValueError('data vazia')
        instante = _parse_data_texto(texto)

    if instante.tzinfo is None:
        instante = instante.replace(tzinfo=_fuso_fabrica())

    return instante.astimezone(_fuso_fabrica())


def classificar_turno(valor):
    """Turno do pagamento derivado do instante: 'T1', 'T2', 'T3' ou None."""
    try:
        instante = instante_fabrica(valor)
    except ValueError:
        return None

    minutos = instante.hour * 60 + instante.minute
    if TURNO_T1_INICIO <= minutos < TURNO_T2_INICIO:
        return 'T1'
    if TURNO_T2_INICIO <= minutos < TURNO_T3_INICIO:
        return 'T2'
    return 'T3'  # 22:00 - 05:19, incluindo a virada do dia


def formatar_data(valor):
    """Data no formato 'dd/mm/aaaa', sempre no dia da fábrica. Vazio se não houver data."""
    try:
        return instante_fabrica(valor).strftime(FORMATO_DATA_EXIBICAO)
    except ValueError:
        # Valor estranho no banco: mostra cru em vez de apagar a informação.
        return '' if valor is None else str(valor).strip()


def formatar_valor(campo, valor):
    """Normaliza um campo do banco para o valor que vai para a célula."""
    if campo in CAMPOS_DATA:
        return formatar_data(valor)
    if valor is None:
        return ''
    if campo in CAMPOS_NUMERICOS:
        return int(valor)
    return str(valor).strip()


def valor_coluna(campo, registro):
    """Valor de uma coluna da planilha para um registro do banco."""
    if campo == 'turno':
        # Coluna derivada: não vem do banco, é calculada a partir de dia_analise.
        return classificar_turno(registro.get('dia_analise')) or ''
    return formatar_valor(campo, registro.get(campo))


def montar_linhas(registros, ids):
    """Monta as linhas na ordem dos IDs recebidos.

    Retorna (linhas, ids_que_nao_existem_no_banco). Cada linha é um dict com os
    rótulos da planilha; não usa pandas (era ~1,3 s de import para montar isso).
    """
    por_id = {}
    for linha in registros:
        try:
            por_id[int(linha['id'])] = linha
        except (KeyError, TypeError, ValueError):
            continue

    linhas = []
    faltantes = []
    for identificador in ids:
        registro = por_id.get(identificador)
        if registro is None:
            faltantes.append(identificador)
            continue
        linhas.append({campo_rotulo: valor_coluna(campo, registro) for campo, campo_rotulo in COLUNAS})

    return linhas, faltantes


# ---------------------------------------------------------------------------
# Estilização e gravação
# ---------------------------------------------------------------------------

def escrever_planilha(linhas, destino):
    """Grava o .xlsx estilizado em `destino` e retorna o caminho."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    planilha = Workbook()
    aba = planilha.active
    aba.title = 'Ordens'

    aba.append(ROTULOS)
    for linha in linhas:
        aba.append([linha[rotulo] if linha[rotulo] != '' else None for rotulo in ROTULOS])

    ultima_linha = max(aba.max_row, 1)
    ultima_coluna = get_column_letter(len(COLUNAS))

    _estilizar_header(aba)
    aba.auto_filter.ref = 'A1:%s%d' % (ultima_coluna, ultima_linha)
    aba.freeze_panes = 'A2'
    _aplicar_bordas(aba, ultima_linha)
    _ajustar_larguras(aba, linhas)

    planilha.save(destino)
    return destino


def _estilizar_header(aba):
    """Fundo azul escuro, texto branco, negrito e centralizado na linha 1."""
    from openpyxl.styles import Alignment, Font, PatternFill

    preenchimento = PatternFill(fill_type='solid', fgColor=COR_HEADER)
    fonte = Font(bold=True, color=COR_TEXTO_HEADER)
    alinhamento = Alignment(horizontal='center', vertical='center')

    for celula in aba[1]:
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = alinhamento
    aba.row_dimensions[1].height = 22


def _aplicar_bordas(aba, ultima_linha):
    """Borda fina em todas as células preenchidas (header + dados)."""
    from openpyxl.styles import Border, Side

    lado = Side(style='thin', color=COR_BORDA)
    borda = Border(left=lado, right=lado, top=lado, bottom=lado)

    for linha in aba.iter_rows(min_row=1, max_row=ultima_linha, max_col=len(COLUNAS)):
        for celula in linha:
            celula.border = borda


def _ajustar_larguras(aba, linhas):
    """Auto-fit: largura pelo maior conteúdo da coluna, com um teto."""
    from openpyxl.utils import get_column_letter

    for indice, rotulo in enumerate(ROTULOS, start=1):
        maior = len(rotulo)
        for linha in linhas:
            valor = str(linha.get(rotulo, '')).strip()
            if valor and len(valor) > maior:
                maior = len(valor)
        aba.column_dimensions[get_column_letter(indice)].width = min(
            LARGURA_MAXIMA_COLUNA, maior + PADDING_LARGURA
        )


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _ler_stdin(stdin):
    """Lê todo o stdin como texto (UTF-8). Sem `stdin`, usa o stdin do processo."""
    if stdin is None:
        stdin = sys.stdin
    if stdin is None:
        return ''
    if hasattr(stdin, 'read'):
        return stdin.read()
    return str(stdin)


def main(argv=None, cliente=None, diretorio=None, stdin=None):
    parser = argparse.ArgumentParser(
        description='Gera a planilha executiva (.xlsx) das ordens de produção.'
    )
    parser.add_argument(
        '--ids',
        help='IDs separados por vírgula (consulta o banco). Sem --ids, lê o JSON do stdin.',
    )

    try:
        argumentos = parser.parse_args(argv)
    except SystemExit as saida:  # argparse já escreveu o uso no stderr
        return int(saida.code or 2)

    try:
        if argumentos.ids is not None:
            ids = parse_ids(argumentos.ids)
            carregar_env()
            registros = buscar_ordens(cliente or criar_cliente(), ids)
        else:
            ids, registros = ler_entrada_stdin(_ler_stdin(stdin))

        if not registros:
            raise ValueError(
                'nenhuma ordem encontrada para exportar (%d id(s) informado(s))' % len(ids)
            )

        linhas, faltantes = montar_linhas(registros, ids)

        if faltantes:
            print(
                'Aviso: %d ID(s) não encontrado(s) e ignorado(s): %s'
                % (len(faltantes), ', '.join(str(i) for i in faltantes)),
                file=sys.stderr,
            )

        destino = os.path.join(
            diretorio or tempfile.gettempdir(),
            'ordens_export_%s.xlsx' % uuid.uuid4().hex[:12],
        )
        escrever_planilha(linhas, destino)
    except Exception as erro:  # noqa: BLE001 - qualquer falha vira erro de exportação
        print('Erro na exportação da planilha: %s' % erro, file=sys.stderr)
        return 1

    print(destino)
    return 0


if __name__ == '__main__':
    # O server.js lê o stdout como UTF-8; garante o mesmo encoding no Windows.
    for fluxo in (sys.stdout, sys.stderr, sys.stdin):
        try:
            fluxo.reconfigure(encoding='utf-8')
        except (AttributeError, ValueError):
            pass
    sys.exit(main())
